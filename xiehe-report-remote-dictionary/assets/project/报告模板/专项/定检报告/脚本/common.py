from pathlib import Path
from copy import deepcopy
import csv,hashlib,json,re,zipfile
from lxml import etree as E
from docx import Document
from docx.table import _Row
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1];PROJECT=next(p for p in ROOT.parents if (p/'报告索引.csv').is_file())
W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
NS={'w':W[1:-1],'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
MODULE_NAMES={'hv':'高压开关柜','lv':'低压开关柜','transformer':'变压器','inverter':'组串式逆变器','secondary':'二次设备','other':'其他设备'}
FIELDS=[
 ('J001','报告年份','year','年','request'),('J002','报告周期','report_period','上半年/下半年/全年','request'),
 ('J003','电站资料编码','station_code','编码','station'),('J004','电站名称','station_name','文本','station'),
 ('J005','定检日期','inspection_date','日期','event'),('J006','定检员签认位','inspector_sign','签认','signature'),
 ('J007','编制人签认位','preparer_sign','签认','signature'),('J008','审核人签认位','reviewer_sign','签认','signature'),('J009','批准人签认位','approver_sign','签认','signature'),
 ('J010','报告出具单位','issuing_org','文本','event'),('J011','并网容量','capacity','MW','station'),('J012','并网电压','voltage','kV','station'),('J013','并网时间','grid_date','日期','station'),
 ('J014','定检参与人员','inspection_people','文本','event'),('J015','本次定检依据','basis','标准及版本','event'),('J016','本次使用仪器','instruments','清单','event'),('J017','安全工器具及其他','safety_tools','清单','event'),('J018','安全措施','safety_measures','文本','event'),('J019','注意事项','precautions','文本','event'),
 ('J020','工作量序号','workload.seq','序号','workload'),('J021','工作量设备名称','workload.device','文本','workload'),('J022','工作量设备型号','workload.model','文本','workload'),('J023','工作量计量单位','workload.unit','单位','workload'),('J024','设备数量','workload.quantity','按计量单位','workload'),('J025','实际定检完成数量','workload.completed','按计量单位','workload'),('J026','定检结果','workload.result','文本','workload'),('J027','工作量备注','workload.notes','文本','workload'),
 ('J028','缺陷序号','defects.seq','序号','defects'),('J029','缺陷设备名称','defects.device','文本','defects'),('J030','缺陷隐患项目','defects.issue','文本','defects'),('J031','缺陷处理结果','defects.treatment','文本','defects'),('J032','缺陷备注','defects.notes','文本','defects'),
 ('J033','遗留问题整改计划','remaining_plan','文本','event'),('J034','定检结论','conclusion','文本','event'),
 ('J035','设备结果序号','results.seq','序号','results'),('J036','设备名称或编号','results.device_id','文本','results'),('J037','设备检查情况','results.finding','文本','results'),('J038','设备定检照片','results.photo','图片','results'),('J039','设备处理结果','results.treatment','文本','results'),
 ('J040','票据类型','tickets.type','文本','tickets'),('J041','票据编号','tickets.number','编号','tickets'),('J042','票据附件','tickets.attachment','文件','tickets')]

def read(p):return json.loads(Path(p).read_text(encoding='utf-8'))
def write(p,data):p=Path(p);p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')
def digest(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def csvwrite(p,rows):
    with Path(p).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
def text(el):return ''.join(el.itertext()) if False else ''.join(el.xpath('.//w:t/text()',namespaces=NS))
def norm(s):return re.sub(r'\s+','',s)
def unique(row):
    out=[]
    for c in row.cells:
        if c._tc not in [x._tc for x in out]:out.append(c)
    return out
def replace_span(p,a,b,value):
    pos=0;done=False
    for r in p.runs:
        end=pos+len(r.text)
        if pos<b and end>a:r.text=r.text[:max(0,a-pos)]+(value if not done else '')+(r.text[b-pos:] if end>b else '');done=True
        pos=end
    if not done:raise ValueError('无法定位替换文字：'+p.text)
def settext(p,value):
    if p.text:replace_span(p,0,len(p.text),value)
    else:p.add_run(value)
def setcell(c,value):
    for el in list(c._tc.xpath('.//w:drawing|.//w:pict|.//w:object')):el.getparent().remove(el)
    settext(c.paragraphs[0],value)
    for p in c.paragraphs[1:]:c._tc.remove(p._p)
def make_proto(table,rows,ids):
    tr=next((deepcopy(r) for r in rows if len(unique(_Row(r,table)))==len(ids)),None)
    if tr is None:raise ValueError('没有匹配列数的原型行：'+str(ids))
    for el in tr.xpath('.//w:vMerge|./w:trPr/w:trHeight'):el.getparent().remove(el)
    for c,fid in zip(unique(_Row(tr,table)),ids):setcell(c,fid)
    return tr
def sdt(tag,children):
    x=E.Element(W+'sdt');pr=E.SubElement(x,W+'sdtPr');E.SubElement(pr,W+'tag').set(W+'val',tag);body=E.SubElement(x,W+'sdtContent')
    for c in children:body.append(c)
    return x
def find_sdt(root,tag):
    xs=root.xpath('.//w:sdt[w:sdtPr/w:tag[@w:val=$tag]]',namespaces=NS,tag=tag)
    if len(xs)!=1:raise ValueError('模块标记缺失或重复：'+tag)
    return xs[0]
def clean_package(source,output,document_xml):
    with zipfile.ZipFile(source) as z:parts={n:z.read(n) for n in z.namelist() if not n.endswith('/')}
    root=E.fromstring(document_xml)
    for obj in root.xpath('//w:drawing|//w:pict|//w:object',namespaces=NS):obj.getparent().remove(obj)
    for el in root.iter():
        for a in ['descr','title']:
            if a in el.attrib:el.set(a,'')
    parts['word/document.xml']=E.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
    deleted=set()
    for n in list(parts):
        if n.startswith(('word/media/','word/embeddings/','customXml/')) or n=='docProps/thumbnail.jpeg':deleted.add(n);del parts[n]
    for n in list(parts):
        if n.endswith('.rels'):
            rel=E.fromstring(parts[n])
            for r in list(rel):
                if any(r.get('Type','').endswith('/'+typ) for typ in ['image','oleObject','customXml']) or r.get('Target','').lstrip('/') in deleted:rel.remove(r)
            parts[n]=E.tostring(rel,xml_declaration=True,encoding='UTF-8',standalone=True)
    ct=E.fromstring(parts['[Content_Types].xml'])
    for el in list(ct):
        if el.get('PartName','').lstrip('/') in deleted:ct.remove(el)
    parts['[Content_Types].xml']=E.tostring(ct,xml_declaration=True,encoding='UTF-8',standalone=True)
    for n in ['docProps/core.xml','docProps/custom.xml']:
        if n in parts:
            el=E.fromstring(parts[n])
            for child in el.iter():
                if len(child)==0:child.text=''
            parts[n]=E.tostring(el,xml_declaration=True,encoding='UTF-8',standalone=True)
    settings=E.fromstring(parts['word/settings.xml']);u=settings.find(W+'updateFields')
    if u is None:u=E.SubElement(settings,W+'updateFields')
    u.set(W+'val','true');parts['word/settings.xml']=E.tostring(settings,xml_declaration=True,encoding='UTF-8',standalone=True)
    Path(output).parent.mkdir(parents=True,exist_ok=True)
    with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as z:
        for n,b in parts.items():z.writestr(n,b)
